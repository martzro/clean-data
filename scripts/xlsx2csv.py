import openpyxl
import csv
import os
import sys

def excel_to_csv(file_name):
    try:
        # Load the Excel workbook
        workbook = openpyxl.load_workbook(file_name, data_only=True)
        
        # Get the base name of the file (without extension)
        base_name, _ = os.path.splitext(file_name)
        
        # Iterate through all sheet names
        for sheet_name in workbook.sheetnames:
            print(f"Processing sheet: {sheet_name}")
            
            # Get the sheet
            sheet = workbook[sheet_name]
            
            # Construct the CSV file name
            csv_file_name = f"{base_name}_{sheet_name}.csv"
            
            # Open a new CSV file for writing
            with open(csv_file_name, mode='w', newline='', encoding='utf-8') as csv_file:
                writer = csv.writer(csv_file)
                
                # Write rows from the Excel sheet to the CSV
                for row in sheet.iter_rows(values_only=True):
                    writer.writerow(row)
            
            print(f"Saved: {csv_file_name}")
        
        print("Conversion completed successfully!")
    
    except Exception as e:
        print(f"An error occurred: {e}")

if __name__ == "__main__":
    # Check if a file name was provided as an argument
    if len(sys.argv) < 2:
        print("Usage: python script.py <excel_file>")
    else:
        file_name = sys.argv[1]
        excel_to_csv(file_name)
